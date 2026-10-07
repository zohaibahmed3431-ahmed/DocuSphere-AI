import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import unittest
import pandas as pd

from src.reconciliation import reconcile, classify_frames


class ReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.bank = pd.DataFrame({
            "Sr No": ["BANK-987", "BANK-988", "BANK-989", "BANK-990"],
            "Transaction Date": ["2026-09-01", "2026-09-02", "2026-09-03", "2026-09-04"],
            "Description": ["ABC PAYMENT", "XYZ PAYMENT", "DUPLICATE", "ONLY BANK"],
            "Debit": [100, 250, 500, 700],
            "Credit": [0, 0, 0, 0],
        })
        self.software = pd.DataFrame({
            "Voucher No": ["SW-42", "SW-43", "SW-44", "SW-45"],
            "Date": ["2026-09-01", "2026-09-02", "2026-09-03", "2026-09-10"],
            "Narration": ["ABC PAYMENT", "XYZ PAYMENT", "DUPLICATE", "ONLY SOFTWARE"],
            "Amount": [100, 250, 500, 700],
            "Type": ["Debit", "Debit", "Debit", "Debit"],
        })

    def test_original_serials_can_differ_and_get_new_common_id(self):
        result = reconcile(self.bank.iloc[:2], self.software.iloc[:2])
        self.assertEqual(len(result["matched"]), 2)
        self.assertEqual(result["matched"].loc[0, "Bank_Original_Serial"], "BANK-987")
        self.assertEqual(result["matched"].loc[0, "Software_Original_Serial"], "SW-42")
        self.assertEqual(result["matched"].loc[0, "Reconciliation_ID"], "REC-000001")

    def test_amount_alone_never_matches(self):
        bank = pd.DataFrame({"Serial": ["B1"], "Amount": [100]})
        software = pd.DataFrame({"Voucher": ["S1"], "Amount": [100]})
        result = reconcile(bank, software)
        self.assertEqual(len(result["matched"]), 0)
        self.assertEqual(len(result["unmatched"]), 2)

    def test_duplicate_ambiguous_rows_stay_unmatched(self):
        bank = pd.DataFrame({"Serial": ["B1"], "Date": ["2026-09-01"], "Amount": [500]})
        software = pd.DataFrame({"Voucher": ["S1", "S2"], "Date": ["2026-09-01", "2026-09-01"], "Amount": [500, 500]})
        result = reconcile(bank, software)
        self.assertEqual(len(result["matched"]), 0)
        self.assertEqual(len(result["unmatched"]), 3)

    def test_date_mismatch_does_not_match(self):
        bank = pd.DataFrame({"Serial": ["B1"], "Date": ["2026-09-01"], "Amount": [100]})
        software = pd.DataFrame({"Voucher": ["S1"], "Date": ["2026-09-10"], "Amount": [100]})
        result = reconcile(bank, software)
        self.assertEqual(len(result["matched"]), 0)

    def test_classification_uses_labels_and_schema(self):
        frames = {"bank_statement.xlsx::Sheet1": self.bank, "accounting_ledger.xlsx::Ledger": self.software}
        bank, software, warnings = classify_frames(frames)
        self.assertIsNotNone(bank)
        self.assertIsNotNone(software)
        self.assertFalse(warnings)


if __name__ == "__main__":
    unittest.main()

class ReconciliationExportTests(unittest.TestCase):
    def test_excel_exports_are_real_xlsx(self):
        from src.reconciliation import matched_excel_bytes, unmatched_excel_bytes
        bank = pd.DataFrame({"Serial": ["B1"], "Date": ["2026-09-01"], "Amount": [10]})
        software = pd.DataFrame({"Voucher": ["S1"], "Date": ["2026-09-01"], "Amount": [10]})
        result = reconcile(bank, software)
        self.assertTrue(matched_excel_bytes(result["matched"]).startswith(b"PK"))
        self.assertTrue(unmatched_excel_bytes(result["bank_unmatched"], result["software_unmatched"]).startswith(b"PK"))

    def test_downloads_are_persisted_in_chat_history_source(self):
        source = (Path(__file__).resolve().parents[1] / "app.py").read_text()
        self.assertIn('item.get("downloads")', source)
        self.assertIn('"downloads":downloads', source)
        self.assertIn('history-download-{i}-{di}', source)

class AdditionalSafetyTests(unittest.TestCase):
    def test_date_mismatch_never_matches_via_description(self):
        bank = pd.DataFrame({"Serial": ["B1"], "Date": ["2026-09-01"], "Amount": [100], "Description": ["foo"]})
        software = pd.DataFrame({"Voucher": ["S1"], "Date": ["2026-09-10"], "Amount": [100], "Narration": ["foo"]})
        result = reconcile(bank, software)
        self.assertEqual(len(result["matched"]), 0)

    def test_direction_mismatch_never_matches(self):
        bank = pd.DataFrame({"Serial": ["B1"], "Date": ["2026-09-01"], "Amount": [100], "Type": ["Debit"]})
        software = pd.DataFrame({"Voucher": ["S1"], "Date": ["2026-09-01"], "Amount": [100], "Type": ["Credit"]})
        result = reconcile(bank, software)
        self.assertEqual(len(result["matched"]), 0)

    def test_download_bytes_round_trip(self):
        from src.reconciliation import matched_excel_bytes, unmatched_excel_bytes
        from openpyxl import load_workbook
        import io
        b = pd.DataFrame({"Serial": ["B1"], "Date": ["2026-09-01"], "Amount": [100], "Description": ["foo"]})
        s = pd.DataFrame({"Voucher": ["S1"], "Date": ["2026-09-01"], "Amount": [100], "Narration": ["foo"]})
        result = reconcile(b, s)
        wb = load_workbook(io.BytesIO(matched_excel_bytes(result["matched"])), read_only=True)
        self.assertIn("Matched", wb.sheetnames)
        wb2 = load_workbook(io.BytesIO(unmatched_excel_bytes(result["bank_unmatched"], result["software_unmatched"])), read_only=True)
        self.assertIn("All_Unmatched", wb2.sheetnames)

    def test_one_day_date_tolerance_can_match_unique_transaction(self):
        bank = pd.DataFrame({"Serial": ["B1"], "Date": ["2026-09-01"], "Amount": [100], "Description": ["foo"]})
        software = pd.DataFrame({"Voucher": ["S1"], "Date": ["2026-09-02"], "Amount": [100], "Narration": ["foo"]})
        result = reconcile(bank, software, date_tolerance_days=1)
        self.assertEqual(len(result["matched"]), 1)

