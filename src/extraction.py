
from __future__ import annotations
import io
from pathlib import Path
from typing import BinaryIO
import pandas as pd
from PIL import Image
import pytesseract
from pypdf import PdfReader

try:
    import pypdfium2 as pdfium
except ImportError:
    pdfium = None
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

TEXT_EXTENSIONS = {
    ".txt",".md",".log",".csv",".py",".java",".cpp",".c",".h",".hpp",".js",".jsx",".ts",".tsx",".html",".css",".scss",".sql",".json",".xml",".yaml",".yml",".toml",".ini",".cfg",".conf",".sh",".bat",".ps1",".php",".go",".rs",".rb",".kt",".kts",".swift",".r",".m",".vue",".svelte",".tex"
}
IMAGE_EXTENSIONS = {".jpg",".jpeg",".png",".webp",".gif",".bmp",".tif",".tiff",".avif"}

def _decode_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")

def extract_pdf(data: bytes) -> list[dict]:
    reader = PdfReader(io.BytesIO(data))
    pages=[]
    needs_ocr=[]
    for number,page in enumerate(reader.pages,1):
        text=page.extract_text() or ""
        pages.append({"text":text,"page":number})
        if not text.strip():
            needs_ocr.append(number)

    # Scanned/image-only PDFs have no text layer. Render those pages and OCR them.
    if needs_ocr and pdfium is not None:
        pdf = pdfium.PdfDocument(io.BytesIO(data))
        for number in needs_ocr:
            page = pdf[number - 1]
            bitmap = page.render(scale=2.0)
            image = bitmap.to_pil()
            text = pytesseract.image_to_string(image)
            pages[number - 1]["text"] = text or ""
            pages[number - 1]["mime_type"] = "image/png"
            buffer = io.BytesIO()
            image.save(buffer, format="PNG")
            pages[number - 1]["image_bytes"] = buffer.getvalue()
    return pages

def extract_docx(data: bytes) -> list[dict]:
    doc=Document(io.BytesIO(data))
    parts=[p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    text="\n".join(parts)
    return [{"text":text,"page":None}] if text else []

def extract_pptx(data: bytes) -> list[dict]:
    prs=Presentation(io.BytesIO(data)); out=[]
    for n,slide in enumerate(prs.slides,1):
        parts=[]
        for shape in slide.shapes:
            if hasattr(shape,"text") and shape.text.strip(): parts.append(shape.text.strip())
        text="\n".join(parts)
        if text: out.append({"text":text,"page":n})
    return out

def extract_xlsx(data: bytes) -> list[dict]:
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    out=[]
    for ws in wb.worksheets:
        rows=[]
        for row in ws.iter_rows(values_only=True):
            vals=["" if v is None else str(v) for v in row]
            if any(v.strip() for v in vals): rows.append(" | ".join(vals))
        if rows: out.append({"text":f"Sheet: {ws.title}\n"+"\n".join(rows),"page":None,"sheet":ws.title})
    return out


def extract_xls(data: bytes) -> list[dict]:
    # Legacy Excel .xls files are handled through pandas/xlrd.
    out=[]
    workbook = pd.ExcelFile(io.BytesIO(data), engine="xlrd")
    for sheet in workbook.sheet_names:
        df = pd.read_excel(io.BytesIO(data), sheet_name=sheet, engine="xlrd")
        text = df.to_csv(index=False)
        if text.strip():
            out.append({"text": f"Sheet: {sheet}\n{text}", "page": None, "sheet": sheet})
    return out

def extract_csv(data: bytes) -> list[dict]:
    bio=io.BytesIO(data)
    try: df=pd.read_csv(bio)
    except Exception:
        bio.seek(0); df=pd.read_csv(bio,encoding="latin1")
    df.columns=[str(c).strip() for c in df.columns]
    # Keep a readable representation for RAG; structured dataframe is handled separately in app.
    text=df.to_csv(index=False)
    return [{"text":text,"page":None,"sheet":None}]

def extract_image(data: bytes) -> list[dict]:
    image=Image.open(io.BytesIO(data))
    text=pytesseract.image_to_string(image)
    return [{"text":text,"page":1,"mime_type":Image.MIME.get(image.format, "image/png"),"image_bytes":data}] if text.strip() else [{"text":"","page":1,"mime_type":Image.MIME.get(image.format, "image/png"),"image_bytes":data}]

def extract_file(uploaded_file: BinaryIO, filename: str) -> list[dict]:
    data=uploaded_file.read()
    suffix=Path(filename).suffix.lower()
    if suffix==".pdf": return extract_pdf(data)
    if suffix==".docx": return extract_docx(data)
    if suffix==".pptx": return extract_pptx(data)
    if suffix==".xlsx": return extract_xlsx(data)
    if suffix==".xls": return extract_xls(data)
    if suffix==".csv": return extract_csv(data)
    if suffix in TEXT_EXTENSIONS:
        text=_decode_text(data); return [{"text":text,"page":None}] if text.strip() else []
    if suffix in IMAGE_EXTENSIONS: return extract_image(data)
    raise ValueError(f"Unsupported file type: {suffix}")
