import io
import os
import re
import json

import streamlit as st
import pandas as pd

from openai import OpenAI

OPENROUTER_BASE_URL = "https://openrouter.ai/api/v1"
OPENROUTER_MODEL = "anthropic/claude-3.5-sonnet"

# ---------------- Setup ----------------

st.set_page_config(
    page_title="Bank Reconciliation AI",
    page_icon="🏦",
    layout="wide",
)

st.markdown(
    """
    <style>
        .stApp {
            background-color: #0F1115;
            color: #E6E6E6;
        }
        .main .block-container {
            max-width: 1100px;
            padding-top: 1.5rem;
            padding-bottom: 6rem;
        }
        h1 {
            font-size: 1.6rem !important;
            font-weight: 700;
        }
        .stChatInput {
            max-width: 720px;
        }
        div[data-testid="stChatInput"] textarea {
            min-height: 44px !important;
            max-height: 44px !important;
            font-size: 0.95rem;
        }
        div[data-testid="stFileUploader"] {
            border: 1px dashed #3A3F4B;
            border-radius: 10px;
            padding: 6px;
            max-width: 720px;
        }
        .stTabs [data-baseweb="tab-list"] {
            gap: 8px;
        }
        .stTabs [data-baseweb="tab"] {
            font-size: 0.9rem;
            padding: 6px 14px;
        }
        .stDataFrame {
            font-size: 0.85rem;
        }
        [data-testid="stMetricValue"] {
            font-size: 1.3rem;
        }
    </style>
    """,
    unsafe_allow_html=True,
)
if "messages" not in st.session_state:
    st.session_state.messages = []

if "uploaded_files" not in st.session_state:
    st.session_state.uploaded_files = []

if "reconciliation" not in st.session_state:
    st.session_state.reconciliation = None

# ---------------- File Reading ----------------

def load_any_file(uploaded_file):
    name = uploaded_file.name.lower()

    if name.endswith((".xlsx", ".xls", ".xlsm")):
        return pd.read_excel(
            io.BytesIO(uploaded_file.getvalue()),
            sheet_name=None,
            header=None,
        )

    if name.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(uploaded_file.getvalue()), header=None, dtype=str)
        return {"CSV": df}

    if name.endswith(".pdf"):
        import pdfplumber

        rows = []
        with pdfplumber.open(io.BytesIO(uploaded_file.getvalue())) as pdf:
            for page in pdf.pages:
                for table in page.extract_tables():
                    rows.extend(table)
                if not page.extract_tables():
                    text = page.extract_text() or ""
                    for line in text.split("\n"):
                        parts = line.split()
                        if parts:
                            rows.append(parts)

        if not rows:
            raise ValueError("PDF se readable data nahi mila.")

        max_cols = max(len(r) for r in rows)
        rows = [r + [""] * (max_cols - len(r)) for r in rows]
        return {"PDF_Data": pd.DataFrame(rows)}

    if name.endswith(".docx"):
        from docx import Document

        document = Document(io.BytesIO(uploaded_file.getvalue()))
        rows = []

        for table in document.tables:
            for row in table.rows:
                rows.append([cell.text.strip() for cell in row.cells])

        if not rows:
            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    rows.append([paragraph.text.strip()])

        if not rows:
            raise ValueError("Word file se readable data nahi mila.")

        return {"Word_Data": pd.DataFrame(rows)}

    if name.endswith((".txt", ".md")):
        text = uploaded_file.getvalue().decode("utf-8", errors="ignore")
        rows = [line.split(",") for line in text.split("\n") if line.strip()]
        return {"Text_Data": pd.DataFrame(rows)}

    raise ValueError("Ye file type abhi supported nahi hai.")


# ---------------- Cleaning ----------------

def clean_amount(value):
    if pd.isna(value):
        return 0.0
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    text = str(value)
    for token in ["PKR", "Rs.", "Rs", "USD", "$", ","]:
        text = text.replace(token, "")
    try:
        return round(float(text.strip()), 2)
    except ValueError:
        return 0.0


def clean_date(value):
    parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)
    return None if pd.isna(parsed) else parsed.date()


def clean_text(value):
    return "" if pd.isna(value) else str(value).strip()


def extract_reference(text):
    text = str(text).upper()
    patterns = [
        r"REF#\s*([A-Z0-9]+)",
        r"STAN\s*\(?([A-Z0-9]+)\)?",
        r"SLIP#:\s*\(?([A-Z0-9]+)\)?",
        r"DOC\s*#\s*([A-Z0-9]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    return ""


def clean_sheet(df):
    df = df.dropna(how="all").reset_index(drop=True)

    header_row = 0
    for i in range(min(40, len(df))):
        row_text = " ".join(str(v).lower() for v in df.iloc[i].values)
        if "date" in row_text and ("debit" in row_text or "credit" in row_text or "amount" in row_text):
            header_row = i
            break

    df = df.iloc[header_row:]
    df.columns = [str(c).strip() for c in df.iloc[0]]
    df = df.iloc[1:].reset_index(drop=True)
    return df


def find_column(df, keywords):
    for column in df.columns:
        name = str(column).lower()
        if any(k in name for k in keywords):
            return column
    return None


def standardize(df, source_type, sheet_name, file_name):
    date_col = find_column(df, ["date"])
    desc_col = find_column(df, ["description", "narration", "particulars"])
    debit_col = find_column(df, ["debit", "withdrawal"])
    credit_col = find_column(df, ["credit", "deposit"])
    amount_col = find_column(df, ["amount"])

    if not date_col or not desc_col:
        raise ValueError(f"{file_name} -> {sheet_name}: Date ya Description column nahi mila.")

    rows = []

    for i, row in df.iterrows():
        date = clean_date(row[date_col])
        description = clean_text(row[desc_col])

        debit = clean_amount(row[debit_col]) if debit_col else 0.0
        credit = clean_amount(row[credit_col]) if credit_col else 0.0
        amount = clean_amount(row[amount_col]) if amount_col else 0.0

        if not date or (debit == 0 and credit == 0 and amount == 0):
            continue

        if "opening balance" in description.lower():
            continue

        if source_type == "Bank Statement":
            direction = "Outflow" if debit > 0 else "Inflow"
            final_amount = debit if debit > 0 else (credit if credit > 0 else amount)
        else:
            direction = "Inflow" if debit > 0 else "Outflow"
            final_amount = debit if debit > 0 else (credit if credit > 0 else amount)

        rows.append({
            "Source": source_type,
            "File": file_name,
            "Sheet": sheet_name,
            "Row": i + 2,
            "Date": date,
            "Description": description,
            "Direction": direction,
            "Amount": final_amount,
            "Reference": extract_reference(description),
            "Status": "Unmatched",
            "Match ID": "",
            "Match Rule": "",
        })

    return pd.DataFrame(rows)


# ---------------- Reconciliation ----------------

def run_reconciliation(bank_df, ledger_df):
    bank_df = bank_df.copy()
    ledger_df = ledger_df.copy()
    counter = 0

    # Rule 1: Exact Reference + Amount + Direction
    for b_idx, bank in bank_df.iterrows():
        if bank_df.at[b_idx, "Status"] == "Matched":
            continue

        ref = str(bank["Reference"]).strip()
        if not ref:
            continue

        candidates = ledger_df[
            (ledger_df["Status"] == "Unmatched")
            & (ledger_df["Reference"].astype(str).str.strip() == ref)
            & (ledger_df["Amount"] == bank["Amount"])
            & (ledger_df["Direction"] == bank["Direction"])
        ]

        if len(candidates) != 1:
            continue

        l_idx = candidates.index[0]
        counter += 1
        match_id = f"REC-{counter:04d}"

        bank_df.at[b_idx, "Status"] = "Matched"
        bank_df.at[b_idx, "Match ID"] = match_id
        bank_df.at[b_idx, "Match Rule"] = "Exact Reference"

        ledger_df.at[l_idx, "Status"] = "Matched"
        ledger_df.at[l_idx, "Match ID"] = match_id
        ledger_df.at[l_idx, "Match Rule"] = "Exact Reference"

    # Rule 2: Exact Amount + Same Date (sirf unique candidate)
    for b_idx, bank in bank_df.iterrows():
        if bank_df.at[b_idx, "Status"] == "Matched":
            continue

        candidates = ledger_df[
            (ledger_df["Status"] == "Unmatched")
            & (ledger_df["Amount"] == bank["Amount"])
            & (ledger_df["Direction"] == bank["Direction"])
            & (ledger_df["Date"] == bank["Date"])
        ]

        if len(candidates) != 1:
            continue

        l_idx = candidates.index[0]
        counter += 1
        match_id = f"REC-{counter:04d}"

        bank_df.at[b_idx, "Status"] = "Matched"
        bank_df.at[b_idx, "Match ID"] = match_id
        bank_df.at[b_idx, "Match Rule"] = "Exact Amount + Date"

        ledger_df.at[l_idx, "Status"] = "Matched"
        ledger_df.at[l_idx, "Match ID"] = match_id
        ledger_df.at[l_idx, "Match Rule"] = "Exact Amount + Date"

    # Rule 3: Exact Amount + Date Window (3 din) — Suggested Match
    for b_idx, bank in bank_df.iterrows():
        if bank_df.at[b_idx, "Status"] == "Matched":
            continue

        bank_date = bank["Date"]

        candidates = ledger_df[
            (ledger_df["Status"] == "Unmatched")
            & (ledger_df["Amount"] == bank["Amount"])
            & (ledger_df["Direction"] == bank["Direction"])
            & (ledger_df["Date"].apply(
                lambda d: abs((d - bank_date).days) <= 3
            ))
        ]

        if len(candidates) != 1:
            continue

        l_idx = candidates.index[0]
        counter += 1
        match_id = f"REC-{counter:04d}"

        bank_df.at[b_idx, "Status"] = "Suggested Match"
        bank_df.at[b_idx, "Match ID"] = match_id
        bank_df.at[b_idx, "Match Rule"] = "Amount + Date Window (3 days)"

        ledger_df.at[l_idx, "Status"] = "Suggested Match"
        ledger_df.at[l_idx, "Match ID"] = match_id
        ledger_df.at[l_idx, "Match Rule"] = "Amount + Date Window (3 days)"

    bank_unmatched = bank_df[bank_df["Status"] == "Unmatched"]
    ledger_unmatched = ledger_df[ledger_df["Status"] == "Unmatched"]

    summary = {
        "Bank Transactions": len(bank_df),
        "Ledger Transactions": len(ledger_df),
        "Matched": len(bank_df) - len(bank_unmatched),
        "Bank Unmatched": len(bank_unmatched),
        "Ledger Unmatched": len(ledger_unmatched),
        "Status": "Needs Review" if len(bank_unmatched) or len(ledger_unmatched) else "Reconciled",
    }

    return bank_df, ledger_df, summary


def auto_reconcile_if_possible():
    if st.session_state.reconciliation:
        return st.session_state.reconciliation

    if not st.session_state.uploaded_files:
        return None

    file_info = st.session_state.uploaded_files[0]
    sheet_names = list(file_info["sheets"].keys())

    if len(sheet_names) < 2:
        return None

    bank_sheet_name = sheet_names[0]
    ledger_sheet_name = sheet_names[1]

    bank_clean = clean_sheet(file_info["sheets"][bank_sheet_name])
    ledger_clean = clean_sheet(file_info["sheets"][ledger_sheet_name])

    bank_df = standardize(
        bank_clean,
        "Bank Statement",
        bank_sheet_name,
        file_info["name"],
    )

    ledger_df = standardize(
        ledger_clean,
        "Company Ledger",
        ledger_sheet_name,
        file_info["name"],
    )

    bank_df, ledger_df, summary = run_reconciliation(bank_df, ledger_df)

    st.session_state.reconciliation = {
        "bank": bank_df,
        "ledger": ledger_df,
        "summary": summary,
    }

    return st.session_state.reconciliation

# ---------------- AI ----------------

def build_context(question):
    parts = []
    parts.append("### Uploaded Files")

    for file_info in st.session_state.uploaded_files:
        parts.append(f"- {file_info['name']} ({file_info['type']})")

    for file_info in st.session_state.uploaded_files:
        for sheet_name, df in file_info["sheets"].items():
            sample = df.head(50)
            parts.append(f"\n### File: {file_info['name']} | Sheet: {sheet_name}")
            parts.append(sample.to_csv(index=False))

    reconciliation = auto_reconcile_if_possible()

    if reconciliation:
        parts.append("\n### Reconciliation Summary")
        parts.append(json.dumps(reconciliation["summary"], indent=2, default=str))

        parts.append("\n### Bank Transactions Sample")
        parts.append(reconciliation["bank"].head(50).to_csv(index=False))

        parts.append("\n### Ledger Transactions Sample")
        parts.append(reconciliation["ledger"].head(50).to_csv(index=False))

    parts.append(f"\n### User Question\n{question}")

    return "\n".join(parts)


def ask_ai(question):
    try:
        api_key = st.secrets["OPENROUTER_API_KEY"]
    except Exception:
        return "⚠️ OPENROUTER_API_KEY Streamlit Secrets mein set nahi hai."

    client = OpenAI(
        api_key=api_key,
        base_url=OPENROUTER_BASE_URL,
    )

    context = build_context(question)

    system_prompt = """
You are a professional Bank Reconciliation AI assistant.

Rules:
1. Answer using the provided file data and reconciliation results.
2. If information is not available, say: "Uploaded files mein is sawal ka jawab nahi mila."
3. Never invent amounts, dates, references, or transactions.
4. For reconciliation questions, use the reconciliation summary and transaction data.
5. For general questions, answer professionally using your own knowledge.
6. Mention file name, sheet name, and row number when answering from file data.
7. Reply in the same language style as the user (Urdu/English mix is fine).
8. Be concise, professional, and accurate.
"""

    response = client.chat.completions.create(
        model=OPENROUTER_MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": context},
        ],
        temperature=0.2,
    )

    return response.choices[0].message.content


# ---------------- Report ----------------

def create_excel_report(reconciliation):
    output = io.BytesIO()

    bank_df = reconciliation["bank"]
    ledger_df = reconciliation["ledger"]
    summary = reconciliation["summary"]

    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        pd.DataFrame(summary.items(), columns=["Metric", "Value"]).to_excel(
            writer, sheet_name="Summary", index=False
        )
        bank_df.to_excel(writer, sheet_name="Bank", index=False)
        ledger_df.to_excel(writer, sheet_name="Ledger", index=False)
        bank_df[bank_df["Status"] == "Unmatched"].to_excel(
            writer, sheet_name="Bank Unmatched", index=False
        )
        ledger_df[ledger_df["Status"] == "Unmatched"].to_excel(
            writer, sheet_name="Ledger Unmatched", index=False
        )

    output.seek(0)
    return output


# ---------------- UI ----------------

st.markdown(
    """
    <div style='display:flex; justify-content:space-between; align-items:center;'>
        <h1 style='margin:0;'>🏦 Bank Reconciliation AI</h1>
    </div>
    <p style='color:#9AA0A6; margin-top:4px;'>
        Upload any file and ask anything. AI will analyze your data and respond professionally.
    </p>
    """,
    unsafe_allow_html=True,
)

if st.button("＋ New Chat"):
    st.session_state.messages = []
    st.session_state.uploaded_files = []
    st.session_state.reconciliation = None
    st.rerun()

for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

uploaded_files = st.file_uploader(
    "Attach files",
    type=["xlsx", "xls", "csv", "pdf", "docx", "txt", "md"],
    accept_multiple_files=True,
    label_visibility="collapsed",
)

if uploaded_files:
    new_added = False

    for uploaded_file in uploaded_files:
        if any(f["name"] == uploaded_file.name for f in st.session_state.uploaded_files):
            continue

        try:
            sheets = load_any_file(uploaded_file)

            if len(sheets) >= 2:
                file_type = "Bank Statement + Company Ledger"
            else:
                text = " ".join(
                    str(v).lower()
                    for df in sheets.values()
                    for v in df.astype(str).fillna("").values.flatten()[:3000]
                )

                if "available balance" in text or "stan" in text:
                    file_type = "Bank Statement"
                elif "general ledger" in text or "doc #" in text:
                    file_type = "Company Ledger"
                else:
                    file_type = "General Document"

            st.session_state.uploaded_files.append({
                "name": uploaded_file.name,
                "type": file_type,
                "sheets": sheets,
            })

            st.session_state.messages.append({
                "role": "assistant",
                "content": f"📄 **{uploaded_file.name}** upload ho gayi.\n\nDetected: **{file_type}**",
            })

            new_added = True

        except Exception as error:
            st.session_state.messages.append({
                "role": "assistant",
                "content": f"⚠️ **{uploaded_file.name}** process nahi ho saki.\n\nError: {error}",
            })

    if new_added:
        st.rerun()

user_input = st.chat_input("Ask anything about your files...")

if user_input:
    st.session_state.messages.append({
        "role": "user",
        "content": user_input,
    })

    with st.chat_message("user"):
        st.markdown(user_input)

    with st.chat_message("assistant"):
        with st.spinner("Analyzing..."):
            reply = ask_ai(user_input)

    st.session_state.messages.append({
        "role": "assistant",
        "content": reply,
    })

    st.rerun()

if st.session_state.reconciliation:
    st.divider()
    st.header("📊 Reconciliation Results")

    reconciliation = st.session_state.reconciliation
    summary = reconciliation["summary"]

    col1, col2, col3, col4 = st.columns(4)

    col1.metric("Bank Transactions", summary["Bank Transactions"])
    col2.metric("Ledger Transactions", summary["Ledger Transactions"])
    col3.metric("Matched", summary["Matched"])
    col4.metric("Status", summary["Status"])

    tab1, tab2, tab3, tab4 = st.tabs([
        "Bank",
        "Ledger",
        "Bank Unmatched",
        "Ledger Unmatched",
    ])

    with tab1:
        st.dataframe(reconciliation["bank"], use_container_width=True)

    with tab2:
        st.dataframe(reconciliation["ledger"], use_container_width=True)

    with tab3:
        st.dataframe(
            reconciliation["bank"][reconciliation["bank"]["Status"] == "Unmatched"],
            use_container_width=True,
        )

    with tab4:
        st.dataframe(
            reconciliation["ledger"][reconciliation["ledger"]["Status"] == "Unmatched"],
            use_container_width=True,
        )

    st.download_button(
        label="⬇️ Download Reconciliation Report",
        data=create_excel_report(reconciliation),
        file_name="Bank_Reconciliation_Report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
