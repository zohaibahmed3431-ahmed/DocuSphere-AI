import io
import pandas as pd

from core.utils import get_file_extension


def load_file(uploaded_file):
    ext = get_file_extension(uploaded_file.name)

    if ext in ("xlsx", "xls", "xlsm"):
        return load_excel(uploaded_file)

    if ext == "csv":
        return load_csv(uploaded_file)

    if ext == "pdf":
        return load_pdf(uploaded_file)

    if ext == "docx":
        return load_docx(uploaded_file)

    if ext in ("txt", "md"):
        return load_text(uploaded_file)

    if ext in ("png", "jpg", "jpeg", "tiff"):
        return load_image(uploaded_file)

    raise ValueError(
        f"Unsupported file type: .{ext}. "
        "Supported: Excel, CSV, PDF, Word, TXT, images."
    )


def load_excel(uploaded_file):
    sheets = pd.read_excel(
        io.BytesIO(uploaded_file.getvalue()),
        sheet_name=None,
        header=None,
    )
    return sheets


def load_csv(uploaded_file):
    df = pd.read_csv(
        io.BytesIO(uploaded_file.getvalue()),
        header=None,
        dtype=str,
    )
    return {"CSV": df}


def load_pdf(uploaded_file):
    try:
        import pdfplumber
    except ImportError:
        raise ValueError("PDF support ke liye 'pdfplumber' install karein.")

    tables = {}
    with pdfplumber.open(io.BytesIO(uploaded_file.getvalue())) as pdf:
        all_rows = []
        for page_number, page in enumerate(pdf.pages, start=1):
            page_tables = page.extract_tables()
            for table in page_tables:
                for row in table:
                    all_rows.append(row)

            if not page_tables:
                text = page.extract_text() or ""
                for line in text.split("\n"):
                    parts = [part.strip() for part in line.split()]
                    if parts:
                        all_rows.append(parts)

    if not all_rows:
        raise ValueError("PDF se koi readable table/text nahi mila.")

    max_cols = max(len(row) for row in all_rows)
    normalized = [
        row + [""] * (max_cols - len(row))
        for row in all_rows
    ]

    df = pd.DataFrame(normalized)
    tables["PDF_Data"] = df
    return tables


def load_docx(uploaded_file):
    try:
        from docx import Document
    except ImportError:
        raise ValueError("Word support ke liye 'python-docx' install karein.")

    document = Document(io.BytesIO(uploaded_file.getvalue()))
    tables = {}

    table_count = 0
    for table in document.tables:
        table_count += 1
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        if rows:
            df = pd.DataFrame(rows)
            tables[f"Table_{table_count}"] = df

    if not tables:
        lines = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                lines.append([paragraph.text.strip()])

        if not lines:
            raise ValueError("Word file se koi readable data nahi mila.")

        df = pd.DataFrame(lines)
        tables["Document_Text"] = df

    return tables


def load_text(uploaded_file):
    text = uploaded_file.getvalue().decode("utf-8", errors="ignore")
    rows = [line.split(",") for line in text.split("\n") if line.strip()]
    df = pd.DataFrame(rows)
    return {"Text_Data": df}


def load_image(uploaded_file):
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        raise ValueError(
            "Image support ke liye 'pytesseract' aur 'Pillow' install karein."
        )

    image = Image.open(io.BytesIO(uploaded_file.getvalue()))
    text = pytesseract.image_to_string(image)

    rows = []
    for line in text.split("\n"):
        parts = [part.strip() for part in line.split()]
        if parts:
            rows.append(parts)

    if not rows:
        raise ValueError("Image se koi readable text nahi mila.")

    max_cols = max(len(row) for row in rows)
    normalized = [row + [""] * (max_cols - len(row)) for row in rows]

    df = pd.DataFrame(normalized)
    return {"OCR_Data": df}